from io import BytesIO

from odoo import models
from odoo.exceptions import UserError


class AgedReceivablesExportMixin(models.AbstractModel):
    _name = "aged.receivables.export.mixin"
    _description = "AR Export Mixin"

    def generate_xlsx_data(self, records):
        """Generate Excel workbook and return as bytes"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            raise UserError(
                "Please install openpyxl library: pip install openpyxl"
            )

        wb = Workbook()
        ws = wb.active
        ws.title = "AR Dashboard"

        # Headers
        headers = [
            "Number",
            "Client",
            "Invoice Date",
            "Due Date",
            "Days Overdue",
            "Total / Received",
        ]

        header_fill = PatternFill(
            start_color="366092", end_color="366092", fill_type="solid"
        )
        header_font = Font(color="FFFFFF", bold=True)

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Data rows
        for row, record in enumerate(records, 2):
            ws.cell(row, 1, record.invoice_name)
            ws.cell(row, 2, record.partner_id.name)
            ws.cell(
                row,
                3,
                (
                    record.invoice_date.strftime("%Y-%m-%d")
                    if record.invoice_date
                    else ""
                ),
            )
            ws.cell(
                row,
                4,
                (
                    record.invoice_due_date.strftime("%Y-%m-%d")
                    if record.invoice_due_date
                    else ""
                ),
            )
            ws.cell(row, 5, record.days_overdue)
            ws.cell(row, 6, record.total_remainder_display)

        ws.cell(row+2, 1, "Summary")
        ws.cell(row+2, 1, "Total Outstanding:")
        ws.cell(row+2, 2, f"{sum(records.mapped('amount_residual'))}")
        ws.cell(row+3, 1, "Total Records:")
        ws.cell(row+3, 2, f"{len(records)}")
        # Adjust column widths
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 15

        # Save to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return output.read()
